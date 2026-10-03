"""Offline contract/regression checks: no network or production database writes."""
import io
import sys
import types
import unittest
from copy import deepcopy
from unittest.mock import patch

# Streamlit is optional for this test runner; exercise UI flow with explicit inputs.
ui=types.ModuleType('streamlit')
sys.modules.setdefault('streamlit',ui)
import pc_connected_research as connected
import pc_document_loader as documents
import pc_source_graph as graph
import pc_research_dossier as dossier
import pc_v07_core as core
from pc_graph_validator import validate_dossier
from pc_v16_research import _save_extra_objects
from pc_intelligence_pipeline import _candidate_research

class Query:
    def __init__(self,db,table):
        self.db=db; self.table=table; self.filters=[]; self.op='select'; self.data=None; self.cap=None; self.conflict=''
    def select(self,*args,**kw): return self
    def eq(self,k,v):
        if self.table=='pc_v10_publication_items' and k=='ingestion_job_id':
            raise ValueError('column pc_v10_publication_items.ingestion_job_id does not exist')
        self.filters.append(lambda x:x.get(k)==v); return self
    def contains(self,k,v): self.filters.append(lambda x:all((x.get(k) or {}).get(a)==b for a,b in v.items())); return self
    def in_(self,k,v): self.filters.append(lambda x:x.get(k) in v); return self
    def order(self,*args,**kw): return self
    def limit(self,n): self.cap=n; return self
    def range(self,a,b): self.cap=b-a+1; return self
    def insert(self,data): self.op='insert'; self.data=data; return self
    def update(self,data): self.op='update'; self.data=data; return self
    def upsert(self,data,**kw): self.op='upsert'; self.data=data; self.conflict=kw.get('on_conflict',''); return self
    def execute(self):
        rows=self.db.rows.setdefault(self.table,[])
        matched=[r for r in rows if all(f(r) for f in self.filters)]
        if self.cap is not None: matched=matched[:self.cap]
        if self.op in ('insert','upsert'):
            matched=[]
            for data in self.data if isinstance(self.data,list) else [self.data]:
                data=deepcopy(data); keys=self.conflict.split(',') if self.conflict else []
                old=next((r for r in rows if keys and all(r.get(k)==data.get(k) for k in keys)),None)
                if old is not None: old.update(data); matched.append(old)
                else:
                    for table,pk in [('pc_documents','document_id'),('pc_ingestion_jobs','ingestion_job_id'),('pc_people','person_id')]:
                        if table==self.table: data.setdefault(pk,str(len(rows)+1))
                    rows.append(data); matched.append(data)
        elif self.op=='update':
            for r in matched: r.update(deepcopy(self.data))
        return types.SimpleNamespace(data=deepcopy(matched),count=len(matched))
class DB:
    def __init__(self,rows=None): self.rows=deepcopy(rows or {})
    def table(self,t): return Query(self,t)


class ResolverDB(DB):
    def rpc(self,name,args):
        def execute():
            table,pk={'pc_resolve_upsert_entity':('pc_entities','entity_id'),
                      'pc_resolve_upsert_asset':('pc_assets','asset_id'),
                      'pc_resolve_upsert_mobile_asset':('pc_mobile_assets','mobile_asset_id')}[name]
            payload=deepcopy(args['p_payload']); rows=self.rows.setdefault(table,[])
            matches=[r for r in rows if (r.get('imo')==payload.get('imo') if payload.get('imo') else r.get('name')==payload.get('name'))]
            if len(matches)>1: return types.SimpleNamespace(data={'status':'REVIEW','reason':'ambiguous'})
            if matches: row=matches[0]
            else:
                row={**payload,pk:table+'_'+str(len(rows)+1)}; rows.append(row)
            return types.SimpleNamespace(data={'status':'OK','canonical_id':row[pk]})
        return types.SimpleNamespace(execute=execute)


def event(url='https://a.test/report',title='Fire at anchorage',day='2026-09-14',ref='A-01',description='18 crew'):
    return {'table':'pc_events','natural_key':title,'payload':{'title':title,'start_date':day,'event_type':'fire',
        'location':'Fujairah','description':description,'metadata':{
        'incident_authority':'UKMTO','incident_reference':ref,'involved_identifiers':['imo:8716306'],
        'intake_source_url':url,'intake_source_key':graph.source_key(url),'research_sources':[url]}}}

class RegressionTests(unittest.TestCase):
    def test_same_event_two_sources_preserves_disagreement(self):
        rows=graph.reconcile_records([event(),event('https://b.test/report',title='Vessel fire',description='19 crew')])
        self.assertEqual(len(rows),1)
        m=rows[0]['payload']['metadata']
        self.assertEqual(len(m['source_observations']),2)
        self.assertEqual(set(m['research_sources']),{'https://a.test/report','https://b.test/report'})
        self.assertTrue(any(x['field']=='description' for x in m['field_conflicts']))
    def test_same_asset_different_incidents_stay_separate(self):
        self.assertEqual(len(graph.reconcile_records([event(),event(ref='A-02')])),2)
    def test_no_identifiers_never_merge_by_title_alone(self):
        a=event(ref=None); b=event('https://b.test/report',ref=None)
        a['payload']['metadata']['involved_identifiers']=[]
        b['payload']['metadata']['involved_identifiers']=[]
        self.assertEqual(len(graph.reconcile_records([a,b])),2)
    def test_one_source_several_events(self):
        rows=graph.reconcile_records([event(),event(ref='A-02',title='Second incident')])
        self.assertEqual(len(rows),2)
        self.assertEqual(len({r['payload']['metadata']['intake_source_key'] for r in rows}),1)
    def test_existing_event_id_is_reused_with_observations(self):
        old=event()['payload']; old['event_id']='EVT_REAL'
        db=DB({'pc_events':[old]})
        rows=graph.bind_existing_events(db,[event('https://b.test/report',title='Vessel fire')])
        self.assertEqual(rows[0]['payload']['event_id'],'EVT_REAL')
        self.assertEqual(len(rows[0]['payload']['metadata']['source_observations']),2)
    def test_duplicate_existing_events_hold(self):
        old=event()['payload']
        with self.assertRaises(ValueError): graph.bind_existing_events(DB({'pc_events':[{**old,'event_id':'E1'},{**old,'event_id':'E2'}]}),[event()])
    def test_same_vessel_imo_merges_sources_and_history(self):
        a={'table':'pc_mobile_assets','payload':{'name':'St Helena','imo':'8716306','metadata':{'identity_history':[{'identifier_value':'MNG Tahiti'}]}}}
        b=deepcopy(a); b['payload']['name']='M/V St Helena'
        self.assertEqual(len(graph.reconcile_records([a,b])),1)
    def test_dossier_rows_do_not_research_again(self):
        rows=[{'target_table':'pc_mobile_assets','staged_record_id':'1','payload':{'metadata':{'dossier_version':'v3'}}}]
        self.assertEqual(_candidate_research(rows,{}),[])
    def test_research_dependencies_do_not_create_stage_rows(self):
        stage={'ingestion_job_id':'j','staged_record_id':'s','payload':{'name':'X'}}
        db=DB({'pc_staged_records':[stage]})
        extra=[{'table':'pc_entities','payload':{'name':'New company'},'evidence_url':'https://a.test'}]
        _save_extra_objects(db,{'task_id':'t'},stage,extra)
        _save_extra_objects(db,{'task_id':'t'},stage,extra)
        self.assertEqual(len(db.rows['pc_staged_records']),1)
        self.assertEqual(len(db.rows['pc_staged_records'][0]['payload']['metadata']['research_dependent_findings']),1)
    def test_replay_company_vessel_and_event_context(self):
        d={'version':'v3','source_url':'https://a.test','graph':{
            'companies':[{'name':'Owner','source_urls':['https://a.test']}],
            'mobile_assets':[{'name':'St Helena','imo':'8716306','source_urls':['https://a.test']}],
            'events':[{'title':'Fire','source_urls':['https://a.test']}],
            'identity_history':[{'asset_name':'St Helena','imo':'8716306','identifier_type':'name','identifier_value':'MNG Tahiti','source_urls':['https://a.test']}],
            'relationships':[{'source_name':'Owner','target_name':'St Helena','relationship':'owner','source_urls':['https://a.test']}],
            'claims':[{'claim':'Crew count disputed'}]}}
        records=dossier.graph_to_core_records(d)
        rows=[{'ingestion_job_id':'j','target_table':r['table'],'payload':r['payload'],'natural_key':r['natural_key']} for r in records]
        plans=connected._validated_replay_plans(DB({'pc_staged_records':rows}),'j')
        self.assertEqual({p['subject_type'] for p in plans},{'company','vessel'})
        self.assertTrue(all(p['source_context']['events'] and p['source_context']['claims'] for p in plans))
        company=next(p for p in plans if p['subject_type']=='company')
        self.assertEqual(company['identity_history'],[])
        self.assertEqual(company['events'],[])
        vessel=next(p for p in plans if p['subject_type']=='vessel')
        self.assertEqual(len(vessel['identity_history']),1)
    def test_event_only_replay_is_reviewable(self):
        e=event(); e['payload']['metadata']['dossier_version']='v3'
        plans=connected._validated_replay_plans(DB({'pc_staged_records':[{'ingestion_job_id':'j','target_table':'pc_events','payload':e['payload']}]}),'j')
        self.assertEqual(plans[0]['subject_type'],'context')
        self.assertEqual(len(plans[0]['events']),1)
    def test_validation_preserves_prior_holds(self):
        d,r=validate_dossier({'graph':{'validator_holds':[{'reason':'research-only object'}]}})
        self.assertEqual(len(d['graph']['validator_holds']),1)
    def test_issuer_without_evidence_is_held(self):
        db=DB(); eid,_=documents._ensure_source_entity(db,{'name':'Unverified'},None,[])
        self.assertIsNone(eid); self.assertNotIn('pc_entities',db.rows)
    def test_issuer_identity_mismatch_is_held(self):
        db=DB(); eid,_=documents._ensure_source_entity(db,{'name':'A'},{'name':'B','source_urls':['https://b.test']},[])
        self.assertIsNone(eid)
    def test_docx_tables_are_extracted(self):
        from docx import Document
        doc=Document(); doc.add_paragraph('Annual report'); doc.add_table(rows=1,cols=1).cell(0,0).text='Revenue 123'
        buf=io.BytesIO(); doc.save(buf)
        upload=types.SimpleNamespace(name='annual.docx',getvalue=lambda:buf.getvalue())
        self.assertIn('Revenue 123',documents._text_from_file(upload))
    def test_company_fleet_invalid_imo_is_held(self):
        with patch.object(connected,'_web',return_value=('evidence',['https://a.test'])),patch.object(connected,'_json_struct',return_value={
            'company':{'name':'A','source_urls':['https://a.test']},'vessels':[{'name':'Fake','imo':'8716307','source_urls':['https://a.test']}]}):
            plan=connected.research_company('key','A')
        self.assertEqual(plan['vessels'],[]); self.assertTrue(plan['research_gaps'])
    def test_company_missing_identity_evidence_stops(self):
        with patch.object(connected,'_web',return_value=('evidence',['https://a.test'])),patch.object(connected,'_json_struct',return_value={'company':{'name':'A'}}):
            with self.assertRaises(ValueError): connected.research_company('key','A')
    def test_document_upload_repeat_reuses_row_and_company_job(self):
        db=DB(); ui.secrets={'OPENAI_API_KEY':'test'}; ui.session_state={}
        upload=types.SimpleNamespace(name='report.txt',getvalue=lambda:b'Annual report. '+b'evidence '*20)
        ui.header=ui.caption=ui.success=ui.warning=ui.error=lambda *a,**k:None
        ui.multiselect=lambda *a,**k:['Trade']; ui.checkbox=lambda *a,**k:True
        ui.file_uploader=lambda *a,**k:[upload]; ui.text_area=lambda *a,**k:'https://a.test/report'
        ui.button=lambda *a,**k:True; ui.dataframe=lambda *a,**k:None
        analysis={'title':'Annual','source_organization':{'name':'A','organization_type':'company'},'published_date':'2026-02-30'}
        with patch.object(documents,'_analyse_document',return_value=analysis) as analyse,patch.object(documents,'_research_org',return_value={'name':'A','source_urls':['https://a.test']}),patch.object(connected,'research_company',return_value={'company':{'name':'A','source_urls':['https://a.test']}}) as research:
            documents.render_document_loader(db); documents.render_document_loader(db)
            self.assertEqual(analyse.call_count,1); self.assertEqual(research.call_count,1)
        self.assertEqual(len(db.rows['pc_documents']),1)
        self.assertEqual(len(db.rows['pc_ingestion_jobs']),1)
        self.assertIsNone(db.rows['pc_documents'][0]['published_date'])
        self.assertEqual(len(db.rows['pc_document_entity_links']),1)

    def test_uncertain_in_batch_events_hold_before_queue(self):
        with self.assertRaises(ValueError):
            graph.require_unambiguous_events([event(ref=None),event('https://b.test',title='Another title',ref=None)])
    def test_uncertain_existing_event_holds(self):
        old=event(ref=None)['payload']; old['event_id']='E1'
        with self.assertRaises(ValueError): graph.bind_existing_events(DB({'pc_events':[old]}),[event(title='Another title')])
    def test_source_investigation_does_not_add_research_only_root(self):
        responses=[{'companies':[{'name':'A'}],'mobile_assets':[],'physical_assets':[],'events':[]},
                   {'research_questions':[]},
                   {'primary_subject':{'name':'A','type':'company'},'companies':[
                       {'name':'A','source_urls':['https://a.test']},
                       {'name':'Unrelated','source_urls':['https://a.test']}]}]
        with patch.object(dossier,'_json_chat',side_effect=responses),patch.object(dossier,'_api',return_value={'output':[{'content':[{'type':'output_text','text':'research'}]}]}):
            result=dossier.investigate_source('A is the issuer','https://a.test','Trade','test')
        self.assertEqual([c['name'] for c in result['graph']['companies']],['A'])
        self.assertEqual(len(result['graph']['validator_holds']),1)
        self.assertEqual(result['source_snapshot_text'],'A is the issuer')
    def test_company_publication_facets_repeat_without_duplicates(self):
        db=ResolverDB(); refs=['https://a.test']
        plan={'subject_name':'A','company':{'name':'A','source_urls':refs},
              'offices':[{'office_name':'HQ','office_type':'headquarters','city':'Abu Dhabi','source_urls':refs}],
              'people':[{'name':'Chief Example','position_title':'CEO','source_urls':refs}],
              'vessels':[{'name':'St Helena','imo':'8716306','operator_name':'A','source_urls':refs}],
              'projects':[{'name':'Terminal Example','asset_type':'terminal','source_urls':refs}],
              'footprint':[{'country':'UAE','activity_type':'logistics','source_urls':refs}],
              'transactions':[],'contracts':[],'related_companies':[],'research_gaps':[]}
        first=connected.publish_company_plan(db,'j',plan); second=connected.publish_company_plan(db,'j',plan)
        self.assertTrue(first['complete']); self.assertTrue(second['complete'])
        for table in ('pc_company_profiles','pc_company_offices','pc_people','pc_company_people_roles','pc_project_details','pc_company_operating_footprint'):
            self.assertEqual(len(db.rows[table]),1,table)
        self.assertEqual(len(db.rows['pc_mobile_assets']),1)
        self.assertEqual(len(db.rows['pc_vessel_identity_history']),2)
    def test_pending_company_relationship_is_held(self):
        db=ResolverDB(); refs=['https://a.test']
        plan={'subject_name':'A','company':{'name':'A','source_urls':refs},
              'related_companies':[{'name':'B','relationship':'owns','status':'pending','source_urls':refs}]}
        report=connected.publish_company_plan(db,'j',plan)
        self.assertFalse(report['complete']); self.assertTrue(report['holds'])
        self.assertNotIn('pc_relationships',db.rows)
    def test_dossier_event_asset_link_is_idempotent(self):
        e=event()['payload']; e['metadata']['dossier_version']='v3'
        db=DB({'pc_staged_records':[{'ingestion_job_id':'j','staged_record_id':'s','target_table':'pc_events','payload':e}],
               'pc_v10_publication_items':[{'staged_record_id':'s','canonical_table':'pc_events','canonical_id':'E1'}],
               'pc_mobile_assets':[{'mobile_asset_id':'M1','name':'St Helena','imo':'8716306'}]})
        connected._publish_dossier_event_links(db,'j'); connected._publish_dossier_event_links(db,'j')
        self.assertEqual(len(db.rows['pc_event_links']),1)
        self.assertEqual(db.rows['pc_event_links'][0]['linked_id'],'M1')

    def test_real_st_helena_dossier_shapes(self):
        from test_graph_validator import run
        run()
    def test_structured_identifiers_link_to_existing_vessel(self):
        e=event()['payload']; e['metadata']['dossier_version']='v3'
        e['metadata']['involved_identifiers']=[{'type':'IMO','value':'8716306'},{'type':'vessel_name','value':'St Helena'}]
        db=DB({'pc_staged_records':[{'ingestion_job_id':'j','staged_record_id':'s','target_table':'pc_events','payload':e}],
               'pc_v10_publication_items':[{'staged_record_id':'s','canonical_table':'pc_events','canonical_id':'E1'}],
               'pc_mobile_assets':[{'mobile_asset_id':'M1','name':'St Helena','imo':'8716306'}]})
        report=connected._publish_dossier_event_links(db,'j')
        self.assertEqual(report['linked'],1); self.assertEqual(report['holds'],[])
    def test_st_helena_replay_reports_source_holds_once(self):
        import json
        from pathlib import Path
        v,r=validate_dossier(json.loads((Path(__file__).parent/'fixtures/st_helena_input.json').read_text()))
        rows=[{'ingestion_job_id':'j','target_table':x['table'],'natural_key':x['natural_key'],'payload':x['payload']} for x in dossier.graph_to_core_records(v)]
        plans=connected._validated_replay_plans(DB({'pc_staged_records':rows}),'j')
        self.assertEqual(len(plans),6)
        self.assertEqual(sum(len(p['validator_holds']) for p in plans),5)
    def test_repeated_validation_keeps_claim_and_partial_date(self):
        import json
        from pathlib import Path
        d=json.loads((Path(__file__).parent/'fixtures/st_helena_input.json').read_text())
        a,_=validate_dossier(d); b,_=validate_dossier(a)
        self.assertEqual(a['graph']['claims'],b['graph']['claims'])
        self.assertEqual(a['graph']['identity_history'],b['graph']['identity_history'])

    def test_shipyard_remains_fixed_infrastructure(self):
        v,_=validate_dossier({'graph':{'physical_assets':[{'name':'Example Shipyard','asset_type':'shipyard','source_urls':['https://a.test']} ]}})
        self.assertEqual(len(v['graph']['physical_assets']),1)

    def test_old_saved_review_refreshes_without_research(self):
        import json
        from pathlib import Path
        d=json.loads((Path(__file__).parent/'fixtures/st_helena_input.json').read_text())
        rows=[{'ingestion_job_id':'j','target_table':x['table'],'natural_key':x['natural_key'],'payload':x['payload']} for x in dossier.graph_to_core_records(d)]
        old={'version':'3.5-source-bounded-replay','status':'review','subjects':[],
             'replayed_without_ai':True,'plans':[{'subject_name':'Ambrey','identity_history':[{'identifier_type':'former_name'}]}]}
        db=DB({'pc_staged_records':rows,'pc_ingestion_jobs':[{'ingestion_job_id':'j','source_scope':{'connected_research':old}}]})
        with patch.object(connected,'research_company',side_effect=AssertionError('Unexpected research call')):
            state=connected.init_job_connected(db,'j')
        self.assertEqual(state['version'],'3.6.0-source-bounded-replay')
        self.assertEqual(len(state['plans']),6)
        company=next(p for p in state['plans'] if p['subject_type']=='company')
        vessel=next(p for p in state['plans'] if p['subject_type']=='vessel')
        self.assertEqual(company['identity_history'],[])
        self.assertEqual(vessel['identity_history'][0]['identifier_type'],'name')
        self.assertEqual(sum(len(p['validator_holds']) for p in state['plans']),5)
        self.assertEqual(len(vessel['claims']),1)
        again=connected.init_job_connected(db,'j')
        self.assertEqual(state['refreshed_at'],again['refreshed_at'])

    def test_same_version_stale_review_refreshes(self):
        expected=[{'subject_type':'company','subject_name':'Ambrey','events':[],'claims':[]}]
        stale={'version':'3.6.0-source-bounded-replay','status':'review','subjects':[],
               'replayed_without_ai':True,'plans':[{'subject_name':'Ambrey','events':[{'title':'incident'}],'claims':[{'claim':'shared'}]}]}
        db=DB({'pc_ingestion_jobs':[{'ingestion_job_id':'j','source_scope':{'connected_research':stale}}]})
        with patch.object(connected,'_validated_replay_plans',return_value=expected),patch.object(connected,'research_company',side_effect=AssertionError('No research allowed')):
            state=connected.init_job_connected(db,'j')
        self.assertEqual(state['plans'],expected)
        self.assertTrue(state['replayed_without_ai'])

    def test_publication_lookup_scopes_through_stage_ids(self):
        db=DB({'pc_v10_publication_items':[
            {'staged_record_id':'S1','canonical_id':'E1','canonical_table':'pc_events'},
            {'staged_record_id':'S2','canonical_id':'E2','canonical_table':'pc_events'}]})
        rows=connected._publications_for_stages(db,[{'staged_record_id':'S1'}])
        self.assertEqual([r['canonical_id'] for r in rows],['E1'])
    def test_company_job_subjects_without_publication_job_column(self):
        db=DB({'pc_staged_records':[{'ingestion_job_id':'j','staged_record_id':'S1','target_table':'pc_entities',
                                    'payload':{'name':'A','entity_type':'company'}}],
               'pc_v10_publication_items':[{'staged_record_id':'S1','canonical_table':'pc_entities','canonical_id':'C1'}]})
        subjects=connected.job_subjects(db,'j')
        self.assertEqual(len(subjects),1)
        self.assertEqual(subjects[0]['canonical_id'],'C1')

class PublisherRegressionTests(unittest.TestCase):
    def stage(self, day='2026-09-14'):
        return {'staged_record_id':'S','ingestion_job_id':'j','target_table':'pc_events',
                'source_record_key':'source','natural_key':'Projectile strike on M/V St. Helena in Strait of Hormuz',
                'payload':{'title':'Projectile strike on M/V St. Helena in Strait of Hormuz','start_date':day,
                           'metadata':{'source_url':'https://a.test/report'}}}
    def database(self, events):
        return DB({'pc_events':events,'pc_v08_trade_content':[{'ingestion_job_id':'j','source_record_key':'source'}]})
    def plan(self, events, day='2026-09-14'):
        import pc_v15_bulk_replay as publisher
        return publisher._plan(self.database(events),'j',[self.stage(day)])
    def test_live_timestamp_matches_existing_event(self):
        old={'event_id':'EVT_PC_97538C73504F6A4F7B60','title':self.stage()['natural_key'],
             'start_date':'2026-09-14T00:00:00+00:00'}
        ready,followers,exceptions,pub=self.plan([old])
        self.assertEqual([(d,c) for _,d,c in ready],[('match_existing',old['event_id'])])
        self.assertEqual(exceptions,[])
    def test_title_date_conflict_is_held(self):
        old={'event_id':'E','title':self.stage()['natural_key'],'start_date':'2026-09-15T00:00:00Z'}
        ready,_,exceptions,_=self.plan([old])
        self.assertEqual(ready,[]); self.assertEqual(len(exceptions),1)
    def test_multiple_same_day_events_are_held(self):
        old={'title':self.stage()['natural_key'],'start_date':'2026-09-14T00:00:00Z'}
        ready,_,exceptions,_=self.plan([{**old,'event_id':'E1'},{**old,'event_id':'E2'}])
        self.assertEqual(ready,[]); self.assertEqual(len(exceptions),1)
    def test_missing_event_date_does_not_match(self):
        old={'event_id':'E','title':self.stage()['natural_key'],'start_date':None}
        ready,_,exceptions,_=self.plan([old],None)
        self.assertEqual(ready,[]); self.assertEqual(len(exceptions),1)
    def test_calendar_day_validation(self):
        for value in ('2026-09','2026-02-30','2026-09-14Tgarbage'):
            self.assertIsNone(graph.event_day(value))
        self.assertEqual(graph.event_day('2026-09-14T23:30:00-04:00'),'2026-09-14')
    def test_partial_job_retries_event_link_and_counts_validator_holds(self):
        state={'status':'published_partial','plans':[{'subject_type':'context','subject_name':'source',
               'validator_holds':[{'reason':'uncertain claim'}]}],'publication_report':{'holds':99}}
        db=DB({'pc_ingestion_jobs':[{'ingestion_job_id':'j','source_scope':{'connected_research':state}}]})
        with patch.object(connected,'_publish_dossier_event_links',return_value={'linked':1,'holds':[]}) as links:
            report=connected.publish_job_connected(db,'j')
        links.assert_called_once(); self.assertEqual(report['event_links']['linked'],1)
        self.assertEqual(report['holds'],1)
    def test_retry_reopens_saved_plans_without_research(self):
        plan={'subject_type':'context','subject_name':'source','validator_holds':[]}
        state={'status':'published_partial','subjects':[],'plans':[plan]}
        db=DB({'pc_ingestion_jobs':[{'ingestion_job_id':'j','source_scope':{'connected_research':state}}]})
        with patch.object(connected,'_validated_replay_plans',return_value=[plan]):
            result=connected.init_job_connected(db,'j',retry=True)
        self.assertEqual(result['status'],'review'); self.assertTrue(result['replayed_without_ai'])

if __name__=='__main__': unittest.main(verbosity=2)
