#!/usr/bin/env python3
"""Validate the deployable P&C Trade System v3.0 Excel-backed package."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
APP = BASE / "app.py"
MANIFEST = BASE / "data_manifest.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sheet_rows(path: Path, sheet: str) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        raw_header = next(iterator)
    except StopIteration:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in raw_header]
    rows: list[dict[str, str]] = []
    for values in iterator:
        record = {
            header: "" if value is None else str(value).strip()
            for header, value in zip(headers, values)
            if header
        }
        if any(record.values()):
            rows.append(record)
    return headers, rows


def duplicate_values(rows: list[dict[str, str]], key: str) -> list[str]:
    values = [row.get(key, "").strip() for row in rows if row.get(key, "").strip()]
    return sorted(value for value, count in Counter(values).items() if count > 1)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", action="store_true", help="Print machine-readable results")
    args = parser.parse_args()

    errors: list[str] = []
    notes: list[str] = []
    stats: dict[str, int | str | dict[str, str]] = {}

    for path in [APP, BASE / "requirements.txt", MANIFEST]:
        if not path.exists():
            errors.append(f"Missing required file: {path.relative_to(BASE)}")
    if not DATA.is_dir():
        errors.append("Missing required directory: data")
    if errors:
        return finish(args.json, errors, notes, stats)

    try:
        compile(APP.read_text(encoding="utf-8"), str(APP), "exec")
        notes.append("app.py syntax")
    except Exception as exc:
        errors.append(f"app.py syntax: {exc}")

    try:
        model = json.loads(MANIFEST.read_text(encoding="utf-8"))
    except Exception as exc:
        errors.append(f"data_manifest.json: {exc}")
        return finish(args.json, errors, notes, stats)

    files_contract = model.get("files", {})
    expected_files = sorted(files_contract)
    actual_files = sorted(path.name for path in DATA.glob("*.xlsx"))
    missing_files = sorted(set(expected_files) - set(actual_files))
    extra_files = sorted(set(actual_files) - set(expected_files))
    if missing_files:
        errors.append("Missing workbooks: " + ", ".join(missing_files))
    if extra_files:
        errors.append("Unexpected workbooks: " + ", ".join(extra_files))

    available: dict[str, set[str]] = {}
    formula_errors: list[str] = []
    total_sheets = 0
    for filename in actual_files:
        path = DATA / filename
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
            available[filename] = set(workbook.sheetnames)
            total_sheets += len(workbook.sheetnames)
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and re.search(
                            r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!|SPILL!|CALC!)",
                            cell.value,
                        ):
                            formula_errors.append(f"{filename}:{worksheet.title}!{cell.coordinate}")
        except Exception as exc:
            errors.append(f"Unreadable workbook {filename}: {exc}")

    expected_sheet_count = 0
    for filename, sheets in files_contract.items():
        expected_sheet_count += len(sheets)
        for sheet in sheets:
            if filename in available and sheet not in available[filename]:
                errors.append(f"Missing sheet: {filename} / {sheet}")
    if formula_errors:
        errors.append("Spreadsheet errors: " + ", ".join(formula_errors[:20]))

    app_text = APP.read_text(encoding="utf-8")
    for filename in expected_files:
        if filename not in app_text:
            errors.append(f"app.py does not reference {filename}")

    companies = sheet_rows(DATA / "01_core_entities.xlsx", "Companies")[1]
    ports = sheet_rows(DATA / "02_maritime.xlsx", "Ports")[1]
    vessels = sheet_rows(DATA / "02_maritime.xlsx", "Vessels")[1]
    systems = sheet_rows(DATA / "11_systems_waterways_governance.xlsx", "Systems")[1]
    events = sheet_rows(DATA / "13_events_hazards.xlsx", "Events")[1]
    sanctions = sheet_rows(DATA / "14_trade_policy_compliance.xlsx", "Sanctions Designations")[1]

    for label, rows, key in [
        ("companies", companies, "Company ID"),
        ("ports", ports, "Port ID"),
        ("vessels", vessels, "Vessel ID"),
        ("systems", systems, "System ID"),
        ("events", events, "Event ID"),
    ]:
        duplicates = duplicate_values(rows, key)
        if duplicates:
            errors.append(f"Duplicate {key}: " + ", ".join(duplicates[:20]))
        stats[label] = len(rows)

    imos = [row.get("IMO", "").strip() for row in vessels if row.get("IMO", "").strip()]
    duplicate_imos = sorted(value for value, count in Counter(imos).items() if count > 1)
    if duplicate_imos:
        errors.append("Duplicate nonblank canonical IMO: " + ", ".join(duplicate_imos[:20]))

    stats.update(
        {
            "sanctions": len(sanctions),
            "unique_imos": len(set(imos)),
            "workbooks": len(actual_files),
            "sheets": total_sheets,
            "model_version": str(model.get("version", "")),
            "hashes": {"app.py": sha256(APP), **{name: sha256(DATA / name) for name in actual_files}},
        }
    )
    notes.extend(
        [
            f"{len(actual_files)} canonical workbooks",
            f"{total_sheets} workbook sheets",
            f"{expected_sheet_count} manifest sheet mappings",
            "spreadsheet error scan",
            "canonical primary-key checks",
            "canonical IMO uniqueness",
        ]
    )
    return finish(args.json, errors, notes, stats)


def finish(as_json: bool, errors: list[str], notes: list[str], stats: dict) -> int:
    result = {"status": "PASS" if not errors else "FAIL", "checks": notes, "errors": errors, "stats": stats}
    if as_json:
        print(json.dumps(result, indent=2, ensure_ascii=False))
    else:
        print("P&C TRADE SYSTEM V3.0 DEPLOYMENT VALIDATION")
        print("=" * 42)
        for note in notes:
            print(f"PASS: {note}")
        for error in errors:
            print(f"FAIL: {error}")
        print(f"\nRESULT: {result['status']}")
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
