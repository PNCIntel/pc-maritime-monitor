from __future__ import annotations
import io, json, os, re, hashlib, urllib.request
from pathlib import Path
from datetime import datetime, timezone
from difflib import SequenceMatcher
import streamlit as st

PRODUCTS = [
    'Trade', 'Security', 'Maritime', 'Aviation', 'Rail', 'Road / Trucking',
    'Energy', 'Ports & Infrastructure', 'Foresight', 'Daily Brief', 'Weekly Intelligence'
]


def _norm(value):
    return ' '.join(re.findall(r'[a-z0-9]+', str(value or '').casefold()))


def _numeric_confidence(value):
    """Only persist finite probabilities; qualitative labels remain in extraction evidence."""
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if 0.0 <= number <= 1.0 else None


def _http_json(endpoint, api_key, payload, timeout=140):
    req = urllib.request.Request(endpoint, data=json.dumps(payload).encode('utf-8'),
        headers={'Authorization': f'Bearer {api_key}', 'Content-Type':'application/json'})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode('utf-8'))


def _text_from_file(uploaded):
    data = uploaded.getvalue(); ext = Path(uploaded.name).suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return '\n'.join((p.extract_text() or '') for p in reader.pages)[:180000]
    if ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return '\n'.join(parts)[:180000]
    if ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data)); parts=[]
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text: parts.append(shape.text)
        return '\n'.join(parts)[:180000]
    if ext in {'.txt','.md'}:
        return data.decode('utf-8', errors='replace')[:180000]
    raise ValueError(f'Unsupported document type: {ext}')


def _analyse_document(text, api_key, products, source_url):
    prompt = f'''You are preparing a document record for the Power & Corridors intelligence database.
Return JSON only. Do not invent missing metadata. Use null or [] where evidence is absent.
Products selected by analyst: {products}
Source URL supplied by analyst: {source_url or 'none'}

Confidence values must be numbers between 0 and 1, or null; never qualitative labels.
Required JSON keys:
title, document_type, published_date (YYYY-MM-DD only if explicitly supported),
source_organization {{name, organization_type, website, confidence}},
authors [strings], summary (100-180 words), analytical_abstract (250-500 words),
key_findings [strings], topics [strings], geographies [strings],
mentioned_entities [{{name, entity_type, role, confidence}}],
why_it_matters, source_notes, research_gaps [strings].

Distinguish issuer/source organisation from organisations merely mentioned.
Corporate, government, regulator and think-tank issuers should be identified when supported.
Do not infer ownership, authorship, publication date, or entity identity from weak context.
DOCUMENT:\n{text[:70000]}'''
    result = _http_json('https://api.openai.com/v1/chat/completions', api_key, {
        'model':'gpt-4.1-mini', 'temperature':0,
        'response_format':{'type':'json_object'},
        'messages':[{'role':'system','content':'Extract source-grounded document metadata and intelligence. Output JSON only.'},
                    {'role':'user','content':prompt}]})
    return json.loads(result['choices'][0]['message']['content'])


def _research_org(name, api_key):
    if not name: return None
    prompt = f'''Verify the identity of the organisation named "{name}" for canonical database resolution.
Find authoritative public evidence where possible. Return concise JSON with name, organization_type,
official_website, aliases, country, and source_urls. Do not guess.'''
    try:
        r=_http_json('https://api.openai.com/v1/responses', api_key, {
          'model':'gpt-4.1-mini','tools':[{'type':'web_search_preview'}], 'input':prompt,'max_output_tokens':1200})
        text=[]
        for item in r.get('output',[]):
            for c in item.get('content',[]):
                if c.get('type') in ('output_text','text') and c.get('text'): text.append(c['text'])
        if not text: return None
        joined='\n'.join(text)
        m=re.search(r'\{.*\}',joined,re.S)
        return json.loads(m.group(0)) if m else {'research_text':joined}
    except Exception:
        return None


def _all_entities(sb, limit=25000):
    out=[]
    for start in range(0,limit,500):
        rows=(sb.table('pc_entities').select('entity_id,name,entity_type,subtype,hq_country,metadata')
              .order('entity_id').range(start,start+499).execute().data or [])
        out.extend(rows)
        if len(rows)<500: break
    return out


def _resolve_entity(name, entities):
    key=_norm(name)
    exact=[e for e in entities if _norm(e.get('name'))==key]
    if len(exact)==1: return ('match',exact[0],[])
    if len(exact)>1: return ('ambiguous',None,exact)
    ranked=[]
    for e in entities:
        other=_norm(e.get('name'))
        if not other: continue
        ratio=SequenceMatcher(None,key,other).ratio()
        if ratio>=0.82: ranked.append((ratio,e))
    ranked.sort(key=lambda x:x[0], reverse=True)
    if ranked: return ('review',None,[e for _,e in ranked[:5]])
    return ('new',None,[])


def _new_entity_id(name, entity_type):
    digest=hashlib.sha256(f'{entity_type}|{_norm(name)}'.encode()).hexdigest()[:20].upper()
    return 'ENTITY_AUTO_'+digest


def _ensure_source_entity(sb, org, research, entities):
    name=(org or {}).get('name')
    if not name: return None, 'No source organisation identified'
    status,match,cands=_resolve_entity(name,entities)
    if status=='match': return match['entity_id'], f"Matched existing: {match['name']}"
    if status in {'ambiguous','review'}:
        labels=', '.join(c['name'] for c in cands[:5])
        return None, 'Needs review: '+labels
    from pc_source_graph import norm
    evidence=(research or {}).get('source_urls') or []
    if not evidence or not norm((research or {}).get('name')) == norm(name):
        return None, 'Issuer held: authoritative identity evidence missing or mismatched'
    otype=((org or {}).get('organization_type') or (research or {}).get('organization_type') or 'organization')
    eid=_new_entity_id(name,otype)
    meta={'created_by':'pc_document_loader_v18','document_source_entity':True}
    if research: meta['identity_research']=research
    payload={'entity_id':eid,'name':name,'entity_type':otype,'metadata':meta}
    payload['metadata']['research_sources']=evidence
    sb.table('pc_entities').upsert(payload,on_conflict='entity_id').execute()
    entities.append(payload)
    return eid, f'Created source organisation: {name}'


def render_document_loader(sb):
    st.header('Load documents')
    st.caption('Reports, annual reports, filings, white papers, think-tank papers and presentations. Documents become searchable evidence linked to canonical entities.')
    key=(st.secrets.get('OPENAI_API_KEY') or st.secrets.get('OPENAI_KEY') or os.getenv('OPENAI_API_KEY') or '')
    products=st.multiselect('Products', PRODUCTS, default=['Trade'])
    connected_research=st.checkbox('Research issuing company and connected operations after saving', value=True,
        help='Uses the same company/vessel research engine as Power Admin; unresolved identities are held rather than guessed.')
    uploads=st.file_uploader('Documents', type=['pdf','docx','pptx','txt','md'], accept_multiple_files=True)
    urls=st.text_area('Source URL(s) — optional, one per document in upload order', height=90)
    st.caption('The issuing company, government body, regulator or think tank is matched to the shared canonical entity registry; a source-backed missing organisation can be created automatically.')
    if not st.button('Analyse & save documents', type='primary', disabled=not uploads): return
    if not key:
        st.error('OPENAI_API_KEY is not configured.'); return
    url_list=[u.strip() for u in urls.splitlines() if u.strip()]
    from pc_connected_research import _public_url
    if len(url_list)>len(uploads) or any(not _public_url(u) for u in url_list):
        st.error('Supply valid public URLs, at most one per document in upload order.'); return
    entities=_all_entities(sb)
    results=[]
    for i,f in enumerate(uploads):
        source_url=url_list[i] if i<len(url_list) else ''
        try:
            digest=hashlib.sha256(f.getvalue()).hexdigest()
            prior=(sb.table('pc_documents').select('*').contains('metadata',{'document_sha256':digest}).limit(2).execute().data or [])
            if len(prior)>1: raise ValueError('Duplicate stored document fingerprints require review')
            existing=prior[0] if prior else None
            analysis=((existing.get('metadata') or {}).get('ai_extraction') if existing else None)
            if not analysis:
                text=_text_from_file(f)
                if len(text.strip())<80: raise ValueError('No usable document text extracted; scanned PDF needs OCR')
                analysis=_analyse_document(text,key,products,source_url)
            org=analysis.get('source_organization') or {}
            research=_research_org(org.get('name'),key) if org.get('name') else None
            source_entity_id,msg=_ensure_source_entity(sb,org,research,entities)
            title=analysis.get('title') or Path(f.name).stem
            from pc_graph_validator import _clean_date
            pubdate=_clean_date(analysis.get('published_date'))
            authors=[str(x).strip() for x in (analysis.get('authors') or []) if str(x).strip()]
            topics=[str(x).strip() for x in (analysis.get('topics') or []) if str(x).strip()]
            geos=[str(x).strip() for x in (analysis.get('geographies') or []) if str(x).strip()]
            summary=str(analysis.get('summary') or '').strip()
            abstract=str(analysis.get('analytical_abstract') or '').strip()
            findings=analysis.get('key_findings') or []
            search='\n'.join([title,summary,abstract,' '.join(authors),' '.join(topics),' '.join(geos),str(org.get('name') or '')])
            meta={'document_sha256':digest,'filename':f.name,'why_it_matters':analysis.get('why_it_matters'),'source_notes':analysis.get('source_notes'),
                  'research_gaps':analysis.get('research_gaps') or [],'source_org_research':research,'ai_extraction':analysis}
            row={'title':title,'document_type':analysis.get('document_type'),'published_date':pubdate,
                 'source_entity_id':source_entity_id,'source_name':org.get('name'),'source_url':source_url or None,
                 'authors':authors,'products':products,'summary':summary,'analytical_abstract':abstract,
                 'key_findings':findings,'topics':topics,'geographies':geos,'search_text':search,'metadata':meta,
                 'updated_at':datetime.now(timezone.utc).isoformat()}
            if existing:
                doc_id=existing['document_id']
                oldmeta=existing.get('metadata') or {}
                row['metadata']={**oldmeta,**meta}
                row['metadata']['source_urls']=list(dict.fromkeys((oldmeta.get('source_urls') or [])+([existing.get('source_url')] if existing.get('source_url') else [])+([source_url] if source_url else [])))
                row['products']=list(dict.fromkeys((existing.get('products') or [])+products))
                if not source_url: row['source_url']=existing.get('source_url')
                if not source_entity_id: row['source_entity_id']=existing.get('source_entity_id')
                sb.table('pc_documents').update(row).eq('document_id',doc_id).execute()
            else:
                inserted=sb.table('pc_documents').insert(row).execute().data or []
                if not inserted: raise RuntimeError('Document insert returned no row')
                doc_id=inserted[0]['document_id']
            connected_job=None
            is_company=any(t in str(org.get('organization_type') or '').casefold()
                           for t in ('company','corporation','business','carrier','operator','commercial'))
            if connected_research and org.get('name') and is_company:
                try:
                    from pc_connected_research import research_company
                    oldjobs=(sb.table('pc_ingestion_jobs').select('ingestion_job_id')
                        .eq('job_type','COMPANY_CONNECTED_RESEARCH')
                        .contains('source_scope',{'document_id':str(doc_id)}).limit(1).execute().data or [])
                    if oldjobs:
                        connected_job=str(oldjobs[0]['ingestion_job_id'])
                    else:
                        plan=research_company(key,org.get('name'),[source_url] if source_url else [])
                        jr=sb.table('pc_ingestion_jobs').insert({
                        'job_type':'COMPANY_CONNECTED_RESEARCH',
                        'title':'Document-connected research: '+str(org.get('name')),
                        'status':'review',
                        'source_scope':{'connected_company_plan':plan,'document_id':str(doc_id),
                                        'document_title':title,'workflow':'document_connected_research_v2'},
                        'stats':{'offices':len(plan.get('offices') or []),'people':len(plan.get('people') or []),
                                 'related_companies':len(plan.get('related_companies') or []),'vessels':len(plan.get('vessels') or []),
                                 'transactions':len(plan.get('transactions') or []),'contracts':len(plan.get('contracts') or []),
                                 'projects':len(plan.get('projects') or []),'research_gaps':len(plan.get('research_gaps') or [])}
                    }).execute().data or []
                        connected_job=str(jr[0]['ingestion_job_id']) if jr else None
                    if connected_job: st.session_state['pc_connected_company_job']=connected_job
                except Exception as _crexc:
                    connected_job='HELD: '+str(_crexc)[:160]
            for a in authors:
                sb.table('pc_document_authors').upsert({'document_id':doc_id,'author_name':a,'metadata':{}},on_conflict='document_id,author_name').execute()
            links=[]
            if source_entity_id: links.append((source_entity_id,'issued_by',1.0))
            for e in analysis.get('mentioned_entities') or []:
                ename=e.get('name') if isinstance(e,dict) else None
                if not ename: continue
                status,match,cands=_resolve_entity(ename,entities)
                if status=='match' and match['entity_id']!=source_entity_id:
                    links.append((match['entity_id'],str(e.get('role') or 'mentions'),e.get('confidence')))
            seen=set()
            for eid,rel,conf in links:
                key2=(eid,rel)
                if key2 in seen: continue
                seen.add(key2)
                sb.table('pc_document_entity_links').upsert({'document_id':doc_id,'entity_id':eid,'relationship':rel,
                    'confidence':_numeric_confidence(conf),'evidence':{'document_title':title,'source_url':source_url or None}},
                    on_conflict='document_id,entity_id,relationship').execute()
            results.append({'Document':title,'Source organisation':org.get('name'),'Canonical source':source_entity_id or 'HELD','Status':msg,'Links':len(seen),'Connected research':connected_job or ('issuer identity only' if connected_research and not is_company else 'not requested'), 'Reused document':bool(existing)})
        except Exception as exc:
            results.append({'Document':f.name,'Status':'ERROR: '+str(exc)[:300]})
    failed=sum(1 for r in results if str(r.get('Status') or '').startswith('ERROR:'))
    if failed: st.warning(f'Saved {len(results)-failed} document(s); {failed} failed. See details below.')
    else: st.success(f'Saved {len(results)} document(s).')
    st.dataframe(results,use_container_width=True,hide_index=True)
