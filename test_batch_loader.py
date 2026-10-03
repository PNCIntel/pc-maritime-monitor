"""20-source downstream integration tests. Extraction and Supabase RPCs are doubles.
URL labels are real publisher URLs; generated graph facts are explicit test data.
No production data or inferred article facts are written by this runner.
"""
import json
import unittest
import types
from pathlib import Path
from copy import deepcopy
from test_loader_regressions import DB, event, connected, graph, dossier, core
import pc_v15_bulk_replay as publisher
from pc_intelligence_pipeline import source_publication_report

URLS=json.loads((Path(__file__).parent/'fixtures/batch_20_urls.json').read_text())

def fixture(i):
    url=URLS[i]
    return {'version':'test-fixture-only','source_url':url,'graph':{
        'companies':[{'name':'Fixture Shared Company','entity_type':'company','country':'United Kingdom','source_urls':[url]}],
        'events':[{'title':'Fixture incident '+str(i),'start_date':'2026-09-'+str(i+1).zfill(2),
                   'event_type':'fire','location':'Fixture port '+str(i),'source_urls':[url]}],
        'claims':[{'claim':'Fixture observation '+str(i),'status':'unconfirmed','source_urls':[url]}]}}

class PublicationDB(DB):
    def __init__(self,rows=None,fail_ids=()):
        super().__init__(rows); self.fail_ids=set(fail_ids); self.calls=[]; self.seen=[]
    def rpc(self,name,args):
        def execute():
            self.calls.append(name)
            if name=='pc_v10_backup_staged':return types.SimpleNamespace(data='BACKUP')
            if name!='pc_v10_publish_approved':raise AssertionError(name)
            ids=args['p_stage_ids']
            if self.fail_ids.intersection(ids):raise ValueError('Fixture publication rejected')
            for sid in ids:
                r=next(x for x in self.rows['pc_staged_records'] if x['staged_record_id']==sid)
                approval=next(x for x in self.rows['pc_v10_approvals'] if x['staged_record_id']==sid)
                self.seen.append(deepcopy(r['payload']))
                table=r['target_table'];pk=publisher.ID_TABLES[table]
                cid=approval.get('canonical_id')
                if not cid:
                    cid='FIXTURE_CANON_'+sid
                    self.rows.setdefault(table,[]).append({**deepcopy(r['payload']),pk:cid})
                self.rows.setdefault('pc_v10_publication_items',[]).append({'staged_record_id':sid,'canonical_table':table,'canonical_id':cid})
            return types.SimpleNamespace(data={'published':len(ids)})
        return types.SimpleNamespace(execute=execute)

class BatchTests(unittest.TestCase):
    def package(self):
        records=[r for i in range(20) for r in dossier.graph_to_core_records(fixture(i),'fixture '+str(i))]
        return graph.reconcile_records(records)
    def staged(self,records):
        return [{'ingestion_job_id':'batch','staged_record_id':'S'+str(i),**r} for i,r in enumerate(core.normalize_rows(records))]
    def database(self,staged,**kw):
        return PublicationDB({'pc_ingestion_jobs':[{'ingestion_job_id':'batch','source_scope':{}}],
           'pc_staged_records':staged,'pc_entities':[{'entity_id':'EXISTING','name':'Fixture Shared Company','entity_type':'company','hq_country':'United Kingdom'}],
           'pc_v08_trade_content':[{'ingestion_job_id':'batch','source_record_key':r['source_record_key']} for r in staged if r['target_table']=='pc_events']},**kw)
    def test_twenty_sources_keep_independent_claims_after_company_merge(self):
        package=self.package();self.assertEqual(len(package),21)
        plans=connected._validated_replay_plans(DB({'pc_staged_records':self.staged(package)}),'batch')
        self.assertEqual(len(plans),20)
        contexts={p['source_context_reference']:p['source_context'] for p in plans}
        self.assertEqual(len(contexts),20)
        for i in range(20):
            context=contexts[graph.source_key(URLS[i])]
            self.assertEqual([c['claim'] for c in context['claims']],['Fixture observation '+str(i)])
            self.assertEqual([e['title'] for e in context['events']],['Fixture incident '+str(i)])
    def test_twenty_sources_publish_match_create_and_retry_without_duplicates(self):
        staged=self.staged(self.package());db=self.database(staged)
        ready,followers,holds,_=publisher._plan(db,'batch',staged)
        self.assertEqual(len(ready),21);self.assertEqual(holds,[])
        results,_,failures=publisher._publish_ready(db,'batch',ready,followers,'fixture')
        self.assertEqual(failures,[]);self.assertEqual(len(db.rows['pc_entities']),1)
        self.assertEqual(len(db.rows['pc_events']),20)
        before=deepcopy(db.rows['pc_events'])
        again,_,holds,pub=publisher._plan(db,'batch',staged)
        self.assertEqual(again,[]);self.assertEqual(holds,[]);self.assertEqual(len(pub),21)
        self.assertEqual(before,db.rows['pc_events'])
        report=source_publication_report(staged,pub,holds,[],{'errors':[{'source':'https://fixture.invalid/unavailable','stage':'fetch','error':'Fixture fetch failure'}]})
        self.assertEqual(len(report),21)
        self.assertTrue(all(r['published']==2 for r in report if r['staged']))
        self.assertEqual(len([r for r in report if r.get('intake_errors')]),1)
    def test_failed_publication_isolated_from_other_sources(self):
        staged=self.staged(self.package());bad=next(r for r in staged if r['target_table']=='pc_events')['staged_record_id']
        db=self.database(staged,fail_ids=[bad]);ready,followers,_,_=publisher._plan(db,'batch',staged)
        _,_,fails=publisher._publish_ready(db,'batch',ready,followers,'fixture')
        self.assertEqual([x['Staged record ID'] for x in fails],[bad])
        self.assertEqual(len(db.rows['pc_events']),19)
    def test_ambiguous_event_does_not_block_other_sources(self):
        a=event(ref=None);b=event(url='https://fixture.invalid/b',ref=None,title='Differently worded incident')
        company={'table':'pc_entities','natural_key':'Fixture Co','payload':{'name':'Fixture Co','entity_type':'company'}}
        prepared=graph.prepare_batch_records(DB(),[a,b,company])
        self.assertEqual(len(prepared),3)
        self.assertTrue(all(publisher._validation_issue({'target_table':'pc_events','payload':r['payload']}) for r in prepared[:2]))
        self.assertIsNone(publisher._validation_issue({'target_table':'pc_entities','payload':prepared[2]['payload']}))
    def test_existing_duplicate_incident_is_held_without_stranding_company(self):
        old=event()['payload']
        prepared=graph.prepare_batch_records(DB({'pc_events':[{**old,'event_id':'E1'},{**old,'event_id':'E2'}]}),[
            event(),{'table':'pc_entities','payload':{'name':'Fixture Company','entity_type':'company'}}])
        self.assertIn('canonical_hold',prepared[0]['payload']['metadata'])
        self.assertEqual(prepared[1]['payload']['name'],'Fixture Company')
    def parse_uploaded(self,upload):
        # Isolate the intake parser from the admin page's auth and Streamlit side effects.
        import ast
        tree=ast.parse((Path(__file__).parent/'pc-power-admin.py').read_text())
        node=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='_parse_uploaded')
        ns={'Path':Path,'json':json}
        exec(compile(ast.Module(body=[node],type_ignores=[]),'intake_parser','exec'),ns)
        return ns['_parse_uploaded'](upload)
    def test_printed_pdf_newsletter_text_and_embedded_links_are_extracted(self):
        import io
        from reportlab.pdfgen import canvas
        buf=io.BytesIO();pdf=canvas.Canvas(buf)
        pdf.drawString(40,740,'Fixture printed newsletter: story one is a port update. Story two is a rail update.')
        pdf.drawString(40,720,'This text is an artificial intake parser test and contains no production claims.')
        pdf.linkURL(URLS[0],(40,700,300,720));pdf.save()
        records,inputs=self.parse_uploaded(types.SimpleNamespace(name='newsletter.pdf',getvalue=lambda:buf.getvalue()))
        self.assertEqual(records,[]);self.assertIn('story one',inputs[0]['text']);self.assertIn('Story two',inputs[0]['text'])
        self.assertIn(URLS[0],inputs[0]['text'])
    def test_word_newsletter_tables_are_extracted(self):
        import io
        from docx import Document
        doc=Document();doc.add_paragraph('Fixture printed newsletter containing two separate stories.')
        doc.add_table(rows=1,cols=1).cell(0,0).text='Fixture rail story and company table'
        buf=io.BytesIO();doc.save(buf)
        records,inputs=self.parse_uploaded(types.SimpleNamespace(name='newsletter.docx',getvalue=lambda:buf.getvalue()))
        self.assertEqual(records,[]);self.assertIn('rail story',inputs[0]['text'])
    def test_twenty_urls_plus_two_newsletters_preserve_twenty_two_sources(self):
        records=self.package()
        for label in ('newsletter.pdf','newsletter.docx'):
            d=fixture(0);d['source_url']=None
            d['graph']['events']=[{'title':'Fixture '+label+' story '+str(i),'start_date':'2026-09-01',
                'event_type':'fire','location':'Fixture newsletter port','source_urls':[URLS[0]]} for i in (1,2)]
            records.extend(dossier.graph_to_core_records(d,label))
        staged=self.staged(graph.reconcile_records(records));db=self.database(staged)
        ready,followers,holds,_=publisher._plan(db,'batch',staged)
        self.assertEqual(holds,[]);self.assertEqual(len(ready),25)
        publisher._publish_ready(db,'batch',ready,followers,'fixture')
        self.assertEqual(len(db.rows['pc_events']),24)
        plans=connected._validated_replay_plans(db,'batch')
        self.assertEqual(len(plans),22)
        self.assertNotIn('pc_documents',db.rows)

    def test_same_vessel_two_sources_preserves_event_identity_and_link_retry(self):
        documents=[]
        for i in (0,1):
            d=fixture(i)
            d['graph']['mobile_assets']=[{'name':'Fixture Vessel','asset_type':'vessel','imo':'8716306','source_urls':[URLS[i]]}]
            d['graph']['events'][0]['involved_identifiers']=['imo:8716306']
            d['graph']['events'][0]['incident_reference']='FIXTURE-INCIDENT'
            d['graph']['events'][0]['incident_authority']='Fixture authority'
            documents.extend(dossier.graph_to_core_records(d))
        staged=self.staged(graph.reconcile_records(documents));db=self.database(staged)
        ready,followers,holds,_=publisher._plan(db,'batch',staged)
        self.assertEqual(holds,[])
        publisher._publish_ready(db,'batch',ready,followers,'fixture')
        self.assertEqual(len(db.rows['pc_mobile_assets']),1);self.assertEqual(len(db.rows['pc_events']),1)
        connected._publish_dossier_event_links(db,'batch');connected._publish_dossier_event_links(db,'batch')
        self.assertEqual(len(db.rows['pc_event_links']),1)
        sources=db.rows['pc_event_links'][0]['metadata']['research_sources']
        self.assertEqual(set(sources),{URLS[0],URLS[1]})
    def test_structured_bad_imo_and_military_unit_are_held(self):
        for table,payload in [('pc_mobile_assets',{'name':'Fixture Vessel','imo':'8716307'}),
                              ('pc_assets',{'name':'Fixture Formation','asset_type':'military_unit'})]:
            self.assertIsNotNone(publisher._validation_issue({'target_table':table,'payload':payload}))

    def held_company(self):
        row=self.staged(dossier.graph_to_core_records(fixture(0)))[0]
        task={'ingestion_job_id':'batch','staged_record_id':row['staged_record_id'],'status':'held',
              'error_text':'Fixture incident claim remains unconfirmed','result':{'original_source_confirmed':True,
               'target_table':'pc_entities','payload':{'name':'Fixture Shared Company','entity_type':'company'},'allowed_evidence':[URLS[0]]}}
        db=self.database([row]);db.rows['pc_v16_research_tasks']=[task]
        db.rows['pc_entities'][0]['metadata']={'canonical_fact':'preserved'}
        return row,task,db
    def test_company_identity_matches_without_publishing_uncertain_claims(self):
        row,task,db=self.held_company();canonical=deepcopy(db.rows['pc_entities'])
        ready,_,holds,_=publisher._plan(db,'batch',[row]);self.assertEqual(holds,[])
        self.assertTrue(ready[0][0]['identity_only_match'])
        publisher._publish_ready(db,'batch',ready,[],'fixture')
        self.assertEqual(db.seen[0]['metadata'],{'canonical_fact':'preserved'})
        self.assertEqual(db.rows['pc_staged_records'][0]['payload'],row['payload'])
        self.assertEqual(db.rows['pc_v16_research_tasks'][0],task)
        self.assertEqual(db.rows['pc_entities'],canonical)
        self.assertNotIn('identity_restore_pending',db.rows['pc_ingestion_jobs'][0]['source_scope'])
    def test_failed_identity_publication_restores_original_proposal(self):
        row,task,db=self.held_company();db.fail_ids={row['staged_record_id']}
        ready,_,_,_=publisher._plan(db,'batch',[row]);_,_,failures=publisher._publish_ready(db,'batch',ready,[],'fixture')
        self.assertEqual(len(failures),1)
        self.assertEqual(db.rows['pc_staged_records'][0]['payload'],row['payload'])
    def test_identity_hold_does_not_authorize_new_entity_or_type_conflict(self):
        row,task,db=self.held_company();db.rows['pc_entities']=[]
        self.assertEqual(publisher._plan(db,'batch',[row])[0],[])
        db.rows['pc_entities']=[{'name':'Fixture Shared Company','entity_type':'government','entity_id':'G'}]
        self.assertEqual(publisher._plan(db,'batch',[row])[0],[])
    def test_interrupted_identity_payload_swap_recovers_saved_dossier(self):
        row,task,db=self.held_company();db.rows['pc_ingestion_jobs'][0]['source_scope']['identity_restore_pending']={row['staged_record_id']:row['payload']}
        db.rows['pc_staged_records'][0]['payload']={'name':'Fixture Shared Company','entity_type':'company'}
        publisher._plan(db,'batch',db.rows['pc_staged_records'])
        self.assertEqual(db.rows['pc_staged_records'][0]['payload'],row['payload'])

if __name__=='__main__':unittest.main(verbosity=2)
